"""
One-time script to create sampledata/ folder with sample input files:
  - functional_specification.docx
  - manual_test_cases.docx
  - expected_output.xlsx
Run from the repo root: python create_sampledata.py
"""
import os
from pathlib import Path

OUTPUT_DIR = Path(__file__).parent / "sampledata"
OUTPUT_DIR.mkdir(exist_ok=True)

# ──────────────────────────────────────────────────────────────────────────────
# Functional Specification (.docx)
# ──────────────────────────────────────────────────────────────────────────────
def create_functional_spec():
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    doc.add_heading("E-Commerce Platform – Functional Specification", 0)

    meta = doc.add_paragraph()
    meta.add_run("Project: ").bold = True
    meta.add_run("Online Shopping Application")
    meta.add_run("\nVersion: ").bold = True
    meta.add_run("2.0")
    meta.add_run("\nDate: ").bold = True
    meta.add_run("July 2025")

    doc.add_heading("1. Overview", 1)
    doc.add_paragraph(
        "The E-Commerce Platform is a web-based application that allows users to browse "
        "and search for products, manage a shopping cart, proceed through checkout, and "
        "view order history. The platform must be accessible, performant, and secure."
    )

    doc.add_heading("2. Functional Requirements", 1)

    features = [
        (
            "FEATURE 1 – User Authentication",
            [
                ("REQ-001", "User Registration",
                 "Users must register with a valid email address and a password of at least 8 characters. "
                 "The system must send a confirmation email with a link that expires in 24 hours."),
                ("REQ-002", "User Login",
                 "Registered users can log in with their email and password. "
                 "Three failed attempts must lock the account for 15 minutes. "
                 "Successful login redirects to the dashboard."),
                ("REQ-003", "Password Reset",
                 "Users can request a password-reset link via email. "
                 "The link must expire after 24 hours and can only be used once."),
                ("REQ-004", "User Logout",
                 "Authenticated users can log out at any time. "
                 "The session must be invalidated server-side immediately."),
            ],
        ),
        (
            "FEATURE 2 – Product Browsing",
            [
                ("REQ-005", "Product Listing",
                 "Display products with thumbnail, name, price, and rating. "
                 "Paginate at 12 items per page. Support sorting by price and popularity."),
                ("REQ-006", "Product Search",
                 "Full-text search across product name and description. "
                 "Results must appear within 2 seconds. Show 'No results found' for zero matches."),
                ("REQ-007", "Product Details",
                 "Each product page shows description, specifications, stock level, "
                 "customer reviews (with ratings), and related products."),
            ],
        ),
        (
            "FEATURE 3 – Shopping Cart",
            [
                ("REQ-008", "Add to Cart",
                 "Authenticated users can add any in-stock product to their cart. "
                 "Cart persists across browser sessions via server-side storage."),
                ("REQ-009", "Update Cart",
                 "Users can change item quantity (1–99) or remove items. "
                 "Cart total must update in real time."),
                ("REQ-010", "Cart Summary",
                 "Display subtotal, applicable taxes, estimated shipping, and grand total."),
            ],
        ),
        (
            "FEATURE 4 – Checkout & Payment",
            [
                ("REQ-011", "Shipping Address",
                 "Users must provide a valid shipping address. "
                 "Previously saved addresses can be selected from a dropdown."),
                ("REQ-012", "Payment Processing",
                 "Support credit/debit card via Stripe. "
                 "Card details must never be stored on the platform's servers."),
                ("REQ-013", "Order Confirmation",
                 "After successful payment, show an order-confirmation page and send a "
                 "confirmation email within 30 seconds."),
            ],
        ),
    ]

    for feature_title, reqs in features:
        doc.add_heading(feature_title, 2)
        for req_id, req_name, req_desc in reqs:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{req_id} – {req_name}: ").bold = True
            p.add_run(req_desc)

    doc.add_heading("3. Non-Functional Requirements", 1)
    nfrs = [
        "Performance: Page load time < 3 seconds on a 4G connection.",
        "Security: HTTPS only. OWASP Top-10 mitigations in place.",
        "Availability: 99.9 % uptime SLA.",
        "Accessibility: WCAG 2.1 AA compliant.",
    ]
    for nfr in nfrs:
        doc.add_paragraph(nfr, style="List Bullet")

    path = OUTPUT_DIR / "functional_specification.docx"
    doc.save(path)
    print(f"Created: {path}")


# ──────────────────────────────────────────────────────────────────────────────
# Manual Test Cases (.docx)
# ──────────────────────────────────────────────────────────────────────────────
def create_manual_test_cases():
    from docx import Document

    doc = Document()
    doc.add_heading("E-Commerce Platform – Manual Test Cases", 0)

    meta = doc.add_paragraph()
    meta.add_run("Project: ").bold = True
    meta.add_run("Online Shopping Application")
    meta.add_run("\nVersion: ").bold = True
    meta.add_run("2.0")
    meta.add_run("\nDate: ").bold = True
    meta.add_run("July 2025")

    test_cases = [
        # (ID, Title, Priority, Module, Preconditions, Steps, Expected Results)
        (
            "TC-001", "Valid User Registration", "High", "Authentication",
            ["Registration page is open", "Email is not already registered"],
            [
                "Navigate to /register",
                "Enter email: newuser@example.com",
                "Enter password: SecurePass@1",
                "Enter confirm password: SecurePass@1",
                "Click the 'Register' button",
            ],
            [
                "User is redirected to a success page",
                "Toast message: 'Registration successful. Please check your email.'",
                "Confirmation email is received within 60 seconds",
            ],
        ),
        (
            "TC-002", "Registration with Duplicate Email", "High", "Authentication",
            ["An account with email duplicate@example.com already exists"],
            [
                "Navigate to /register",
                "Enter email: duplicate@example.com",
                "Enter password: SecurePass@1",
                "Click 'Register'",
            ],
            [
                "Error message: 'An account with this email already exists.'",
                "User remains on the registration page",
                "No confirmation email is sent",
            ],
        ),
        (
            "TC-003", "Valid Login", "High", "Authentication",
            ["User account exists and email is confirmed"],
            [
                "Navigate to /login",
                "Enter email: testuser@example.com",
                "Enter password: SecurePass@1",
                "Click 'Login'",
            ],
            [
                "User is redirected to /dashboard",
                "Welcome message displays the user's name",
                "Navigation bar shows 'Logout' link",
            ],
        ),
        (
            "TC-004", "Login with Wrong Password", "High", "Authentication",
            ["User account exists"],
            [
                "Navigate to /login",
                "Enter email: testuser@example.com",
                "Enter password: WrongPassword",
                "Click 'Login'",
            ],
            [
                "Error message: 'Invalid email or password.'",
                "User remains on /login",
                "Password field is cleared",
            ],
        ),
        (
            "TC-005", "Product Search – Positive", "High", "Product Browsing",
            ["At least one product named 'Laptop' exists in the catalogue"],
            [
                "Navigate to the home page",
                "Type 'Laptop' in the search bar",
                "Press Enter or click the search icon",
            ],
            [
                "Search results page loads within 2 seconds",
                "Results contain at least one product with 'Laptop' in its name",
                "Result count is displayed",
            ],
        ),
        (
            "TC-006", "Add Product to Cart", "High", "Shopping Cart",
            ["User is logged in", "Product 'Laptop Pro 15' is in stock"],
            [
                "Navigate to the product detail page of 'Laptop Pro 15'",
                "Set quantity to 2",
                "Click 'Add to Cart'",
            ],
            [
                "Success notification: 'Item added to cart.'",
                "Cart icon counter updates to reflect new item count",
                "Cart contains 2 units of 'Laptop Pro 15'",
            ],
        ),
        (
            "TC-007", "Remove Item from Cart", "Medium", "Shopping Cart",
            ["User is logged in", "Cart contains at least one item"],
            [
                "Navigate to /cart",
                "Click the 'Remove' icon next to the first item",
                "Confirm removal in the dialog",
            ],
            [
                "Item is removed from the cart",
                "Cart total is recalculated immediately",
                "If cart is empty, display 'Your cart is empty.' message",
            ],
        ),
        (
            "TC-008", "Successful Checkout", "High", "Checkout",
            ["User is logged in", "Cart has at least one item", "Valid payment details available"],
            [
                "Navigate to /cart and click 'Proceed to Checkout'",
                "Enter shipping address",
                "Enter card number: 4242 4242 4242 4242, Exp: 12/27, CVV: 123",
                "Click 'Place Order'",
            ],
            [
                "Order confirmation page is displayed with order ID",
                "Confirmation email arrives within 30 seconds",
                "Cart is emptied",
            ],
        ),
        (
            "TC-009", "Password Reset Flow", "Medium", "Authentication",
            ["User account exists with email reset@example.com"],
            [
                "Navigate to /forgot-password",
                "Enter email: reset@example.com",
                "Click 'Send Reset Link'",
                "Open the email and click the reset link",
                "Enter new password: NewSecure@99",
                "Confirm new password: NewSecure@99",
                "Click 'Reset Password'",
            ],
            [
                "Success message: 'Password reset successfully.'",
                "User can log in with the new password",
                "Old password no longer works",
            ],
        ),
        (
            "TC-010", "Account Lockout After 3 Failed Logins", "High", "Authentication",
            ["User account exists"],
            [
                "Attempt login with wrong password three times in a row",
                "On the 4th attempt, try correct credentials",
            ],
            [
                "After 3rd failure: 'Account locked for 15 minutes.'",
                "4th attempt (even with correct password) shows same lockout message",
                "Account unlocks after 15 minutes",
            ],
        ),
    ]

    for tc in test_cases:
        tc_id, title, priority, module, preconditions, steps, expected = tc
        doc.add_heading(f"{tc_id}: {title}", 2)

        info = doc.add_table(rows=3, cols=2)
        info.style = "Table Grid"
        info.cell(0, 0).text = "Priority"
        info.cell(0, 1).text = priority
        info.cell(1, 0).text = "Module"
        info.cell(1, 1).text = module

        doc.add_heading("Preconditions", 3)
        for pre in preconditions:
            doc.add_paragraph(pre, style="List Bullet")

        doc.add_heading("Test Steps", 3)
        for i, step in enumerate(steps, 1):
            doc.add_paragraph(f"{i}. {step}")

        doc.add_heading("Expected Results", 3)
        for exp in expected:
            doc.add_paragraph(exp, style="List Bullet")

        doc.add_paragraph()  # spacer

    path = OUTPUT_DIR / "manual_test_cases.docx"
    doc.save(path)
    print(f"Created: {path}")


# ──────────────────────────────────────────────────────────────────────────────
# Expected Output (.xlsx)
# ──────────────────────────────────────────────────────────────────────────────
def create_expected_output():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()

    # ── Sheet 1: Expected HTTP Responses ──────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Expected HTTP Responses"

    headers1 = ["Test Case ID", "Scenario", "HTTP Method", "Endpoint",
                 "Expected Status", "Expected Response Body (key fields)",
                 "Assertion / Validation"]
    ws1.append(headers1)

    rows1 = [
        ["TC-001", "Valid Registration", "POST", "/api/auth/register", "201 Created",
         '{"status":"success","message":"Registration successful. Please check your email.","user_id":"<uuid>"}',
         "status == 'success'; user_id is a valid UUID; confirmation email sent"],
        ["TC-002", "Duplicate Email", "POST", "/api/auth/register", "409 Conflict",
         '{"status":"error","message":"An account with this email already exists.","field":"email"}',
         "status == 'error'; message contains 'already exists'"],
        ["TC-003", "Valid Login", "POST", "/api/auth/login", "200 OK",
         '{"status":"success","token":"<jwt>","user":{"id":"<uuid>","email":"testuser@example.com"}}',
         "token is a non-empty JWT string; user.email matches input"],
        ["TC-004", "Wrong Password", "POST", "/api/auth/login", "401 Unauthorized",
         '{"status":"error","message":"Invalid email or password."}',
         "status == 'error'; HTTP status is 401"],
        ["TC-005", "Product Search", "GET", "/api/products?q=Laptop", "200 OK",
         '{"results":[...],"total":N,"page":1}',
         "results is a non-empty array; each item has id, name, price, rating"],
        ["TC-006", "Add to Cart", "POST", "/api/cart/items", "200 OK",
         '{"cart_id":"<uuid>","item_count":N,"subtotal":N.NN}',
         "item_count > 0; subtotal is a positive number"],
        ["TC-007", "Remove from Cart", "DELETE", "/api/cart/items/{id}", "200 OK",
         '{"message":"Item removed","item_count":N}',
         "item_count decremented by 1; message == 'Item removed'"],
        ["TC-008", "Successful Checkout", "POST", "/api/orders", "201 Created",
         '{"order_id":"<uuid>","status":"confirmed","total":N.NN}',
         "order_id is a valid UUID; status == 'confirmed'; confirmation email sent within 30 s"],
        ["TC-009", "Password Reset", "POST", "/api/auth/reset-password", "200 OK",
         '{"status":"success","message":"Password reset successfully."}',
         "Old password rejected; new password accepted on next login"],
        ["TC-010", "Account Lockout", "POST", "/api/auth/login", "423 Locked",
         '{"status":"error","message":"Account locked for 15 minutes.","locked_until":"<iso-datetime>"}',
         "HTTP 423; locked_until is a valid ISO datetime ~15 min in the future"],
    ]
    for row in rows1:
        ws1.append(row)

    # Style header row
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2563EB")
    for cell in ws1[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    ws1.column_dimensions["A"].width = 12
    ws1.column_dimensions["B"].width = 28
    ws1.column_dimensions["C"].width = 12
    ws1.column_dimensions["D"].width = 30
    ws1.column_dimensions["E"].width = 18
    ws1.column_dimensions["F"].width = 55
    ws1.column_dimensions["G"].width = 55

    # ── Sheet 2: UI Assertions ────────────────────────────────────────────────
    ws2 = wb.create_sheet("UI Assertions")
    headers2 = ["Test Case ID", "UI Element", "Expected State / Text", "Assertion Type"]
    ws2.append(headers2)

    rows2 = [
        ["TC-001", "Success toast", "Registration successful. Please check your email.", "Text contains"],
        ["TC-002", "Email field error", "An account with this email already exists.", "Text equals"],
        ["TC-003", "Dashboard heading", "Welcome back, Test User", "Text starts with 'Welcome'"],
        ["TC-003", "Navbar logout link", "Logout", "Element visible"],
        ["TC-004", "Login error banner", "Invalid email or password.", "Text equals"],
        ["TC-004", "Password field", "(empty)", "Field value is empty"],
        ["TC-005", "Search result count", ">= 1 result", "Count > 0"],
        ["TC-006", "Cart counter badge", "Updated count", "Incremented by 1"],
        ["TC-007", "Empty cart message", "Your cart is empty.", "Text equals (conditional)"],
        ["TC-008", "Order confirmation heading", "Order Confirmed!", "Text equals"],
        ["TC-008", "Order ID display", "Non-empty UUID", "Matches UUID regex"],
        ["TC-009", "Reset success message", "Password reset successfully.", "Text equals"],
        ["TC-010", "Lockout warning banner", "Account locked for 15 minutes.", "Text contains"],
    ]
    for row in rows2:
        ws2.append(row)

    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 28
    ws2.column_dimensions["C"].width = 48
    ws2.column_dimensions["D"].width = 28

    # ── Sheet 3: Acceptance Criteria Summary ──────────────────────────────────
    ws3 = wb.create_sheet("Acceptance Criteria")
    headers3 = ["Feature", "Requirement ID", "Acceptance Criterion", "Priority", "Automated?"]
    ws3.append(headers3)

    rows3 = [
        ["Authentication", "REQ-001", "New user can register and receive a confirmation email", "High", "Yes"],
        ["Authentication", "REQ-001", "Duplicate email returns a 409 error", "High", "Yes"],
        ["Authentication", "REQ-002", "Valid credentials grant access to dashboard", "High", "Yes"],
        ["Authentication", "REQ-002", "Invalid credentials show error, do not grant access", "High", "Yes"],
        ["Authentication", "REQ-002", "3 failed attempts lock the account for 15 min", "High", "Yes"],
        ["Authentication", "REQ-003", "Password reset link is sent and works within 24 h", "Medium", "Yes"],
        ["Product Browsing", "REQ-005", "Product list loads within 3 seconds", "High", "Yes"],
        ["Product Browsing", "REQ-006", "Search returns relevant results within 2 seconds", "High", "Yes"],
        ["Product Browsing", "REQ-006", "No-match search shows 'No results found'", "Medium", "Yes"],
        ["Shopping Cart", "REQ-008", "In-stock product can be added to cart", "High", "Yes"],
        ["Shopping Cart", "REQ-009", "Item quantity can be updated; total recalculates", "Medium", "Yes"],
        ["Shopping Cart", "REQ-009", "Item can be removed; cart updates", "Medium", "Yes"],
        ["Checkout", "REQ-011", "Valid shipping address is accepted", "High", "Yes"],
        ["Checkout", "REQ-012", "Payment succeeds with test card 4242 4242 4242 4242", "High", "Yes"],
        ["Checkout", "REQ-013", "Order confirmation page and email are generated", "High", "Yes"],
    ]
    for row in rows3:
        ws3.append(row)

    for cell in ws3[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    ws3.column_dimensions["A"].width = 22
    ws3.column_dimensions["B"].width = 16
    ws3.column_dimensions["C"].width = 58
    ws3.column_dimensions["D"].width = 12
    ws3.column_dimensions["E"].width = 14

    path = OUTPUT_DIR / "expected_output.xlsx"
    wb.save(path)
    print(f"Created: {path}")


# ──────────────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    print(f"Creating sample data in: {OUTPUT_DIR}")
    create_functional_spec()
    create_manual_test_cases()
    create_expected_output()
    print("Done. sampledata/ folder is ready.")
