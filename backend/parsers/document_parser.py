from docx import Document
from typing import Dict, List, Any
from loguru import logger
import json

try:
    import pandas as pd
    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False
    logger.warning("pandas not installed. Excel parsing will be unavailable.")


class DocumentParser:
    def __init__(self):
        self.logger = logger

    def parse_docx(self, file_path: str) -> Dict[str, Any]:
        try:
            doc = Document(file_path)
            
            sections = []
            tables = []
            headings = []
            paragraphs = []

            for para in doc.paragraphs:
                if para.style.name.startswith('Heading'):
                    headings.append(para.text)
                else:
                    paragraphs.append(para.text)

            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text for cell in row.cells]
                    table_data.append(row_data)
                tables.append(table_data)

            self.logger.info(f"Successfully parsed DOCX: {file_path}")
            
            return {
                "sections": sections,
                "tables": tables,
                "headings": headings,
                "paragraphs": paragraphs
            }
        except Exception as e:
            self.logger.error(f"Error parsing DOCX {file_path}: {str(e)}")
            raise

    def parse_excel(self, file_path: str) -> Dict[str, Any]:
        if not PANDAS_AVAILABLE:
            self.logger.warning("pandas not available. Returning empty Excel data.")
            return {"sheets": {}, "sheet_names": [], "content": "", "paragraphs": [], "lines": [], "headings": [], "tables": []}
        
        try:
            excel_file = pd.ExcelFile(file_path)
            sheets = {}
            text_parts = []
            tables = []

            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name, dtype=str)
                df = df.fillna("")
                records = df.to_dict(orient='records')
                sheets[sheet_name] = records

                text_parts.append(f"Sheet: {sheet_name}")

                # Build table for structured access
                if not df.empty:
                    header_row = list(df.columns)
                    data_rows = [list(row) for row in df.values]
                    tables.append([header_row] + data_rows)

                # Build text representation for LLM agents
                for record in records:
                    row_text = " | ".join(
                        f"{k}: {v}" for k, v in record.items() if str(v).strip()
                    )
                    if row_text.strip():
                        text_parts.append(row_text)

            content = "\n".join(text_parts)
            self.logger.info(f"Successfully parsed Excel: {file_path}")

            return {
                "sheets": sheets,
                "sheet_names": excel_file.sheet_names,
                "content": content,
                "paragraphs": text_parts,
                "lines": text_parts,
                "headings": [f"Sheet: {s}" for s in excel_file.sheet_names],
                "tables": tables,
            }
        except Exception as e:
            self.logger.error(f"Error parsing Excel {file_path}: {str(e)}")
            raise

    def parse_document(self, file_path: str, document_type: str) -> Dict[str, Any]:
        if file_path.endswith('.docx'):
            return self.parse_docx(file_path)
        elif file_path.endswith('.xlsx') or file_path.endswith('.xls'):
            return self.parse_excel(file_path)
        elif file_path.endswith('.txt'):
            return self.parse_text(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_path}")
    
    def parse_text(self, file_path: str) -> Dict[str, Any]:
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            
            paragraphs = [p.strip() for p in content.split('\n\n') if p.strip()]
            lines = [line.strip() for line in content.split('\n') if line.strip()]
            
            self.logger.info(f"Successfully parsed TXT: {file_path}")
            
            return {
                "paragraphs": paragraphs,
                "lines": lines,
                "headings": [],
                "tables": [],
                "content": content
            }
        except Exception as e:
            self.logger.error(f"Error parsing TXT {file_path}: {str(e)}")
            raise
